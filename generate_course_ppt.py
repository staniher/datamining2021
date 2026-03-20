
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def add_slide(prs, title, points):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title

    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = points[0]

    for point in points[1:]:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0

def create_presentation():
    prs = Presentation()

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Python pour les Ingénieurs & Architectes (2026)"
    subtitle.text = "Cours pour Débutants en Classes Préparatoires\nGenie Info, Civil, Électrique & Architecture"

    # Slide 2: Introduction
    add_slide(prs, "1. Pourquoi Python en 2026 ?", [
        "Le langage n°1 mondial : Polyvalence absolue.",
        "Intégration Native de l'IA : Copilotes et automatisation.",
        "Écosystème Mature : De la simulation physique au rendu 3D.",
        "Standard de l'industrie : Interopérabilité entre logiciels métiers."
    ])

    # Slide 3: Les Bases Fondamentales
    add_slide(prs, "2. Fondations : Variables et Contrôle", [
        "Variables typées dynamiquement (mais typage suggéré en 2026).",
        "Structures de contrôle : if, elif, else.",
        "Boucles modernes : for in, while et list comprehensions.",
        "Exemple : calcul d'une charge ou d'une tension simple."
    ])

    # Slide 4: Structures de Données
    add_slide(prs, "3. Organiser l'Information", [
        "Listes : Stockage de séries de mesures ou coordonnées.",
        "Dictionnaires : Modélisation d'objets (ex: caractéristiques d'un matériau).",
        "Tuples : Données immuables (coordonnées GPS, constantes physiques).",
        "Sets : Gestion de collections uniques."
    ])

    # Slide 5: Fonctions et Modularité
    add_slide(prs, "4. Automatisation et Réutilisation", [
        "Définition de fonctions (def) : Factoriser son code.",
        "Modules et Packages : Importer la puissance de la communauté.",
        "Gestion des erreurs (try/except) : Pour des simulations robustes.",
        "Documentation automatique via Docstrings."
    ])

    # Slide 6: Spécialité : Génie Informatique
    add_slide(prs, "5. Focus : Génie Informatique", [
        "Développement d'Agents IA locaux.",
        "Analyse de données massives (Polars/Pandas 2026).",
        "Cybersécurité et Automatisation des réseaux.",
        "Full-stack Python avec frameworks ultra-rapides."
    ])

    # Slide 7: Spécialité : Génie Civil
    add_slide(prs, "6. Focus : Génie Civil", [
        "Calcul de structures et éléments finis (SciPy).",
        "Analyse de données géotechniques.",
        "Automatisation de rapports d'expertise.",
        "Lien avec les SIG (Systèmes d'Information Géographique)."
    ])

    # Slide 8: Spécialité : Génie Électrique
    add_slide(prs, "7. Focus : Génie Électrique", [
        "Traitement du signal et filtres numériques.",
        "Simulation de circuits de puissance.",
        "Interfaçage avec systèmes embarqués (MicroPython).",
        "Analyse de consommation énergétique et Smart Grids."
    ])

    # Slide 9: Spécialité : Architecture
    add_slide(prs, "8. Focus : Architecture", [
        "Conception paramétrique (Scripts pour Rhino/Grasshopper).",
        "Analyse environnementale et thermique (Ladybug tools).",
        "Génération de formes via algorithmes évolutionnaires.",
        "Extraction de données BIM (Building Information Modeling)."
    ])

    # Slide 10: Conclusion
    add_slide(prs, "Conclusion : Vers la Maîtrise", [
        "Apprendre par la pratique : Les projets métiers.",
        "Utiliser l'IA comme levier, non comme béquille.",
        "La programmation est le nouveau dessin technique.",
        "Prochaine étape : Votre premier script de calcul !"
    ])

    prs.save('Cours_Python_Prepa_2026.pptx')
    print("PowerPoint généré avec succès : Cours_Python_Prepa_2026.pptx")

if __name__ == "__main__":
    create_presentation()
